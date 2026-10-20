"""Export pipeline CSV results to one formatted XLSX workbook.

pandas and openpyxl are imported inside export_workbook so the TUI does not pay
their import cost at startup.
"""

from __future__ import annotations

from pathlib import Path

from tools import ToolError

SHEETS = (
    ('combined', 'combined_384_plate.csv'),
    ('pivot', 'combined_384_plate_pivot.csv'),
    ('metadata', 'metadata.csv'),
    ('plate1', 'plate1.csv'),
    ('plate2', 'plate2.csv'),
    ('plate3', 'plate3.csv'),
    ('plate4', 'plate4.csv'),
)


def _first_line(exc):
    """Return the first line of an exception's root context message."""
    while getattr(exc, '__context__', None) is not None and exc.__context__ is not exc:
        exc = exc.__context__
    lines = str(exc).splitlines()
    return lines[0] if lines else type(exc).__name__


def export_workbook(result_dir) -> Path:
    """Write available CSV files from a result folder into one XLSX."""
    result_dir = Path(result_dir)
    if not result_dir.is_dir():
        raise ToolError(f'folder not found: {result_dir}')

    available = [
        (sheet, result_dir / filename)
        for sheet, filename in SHEETS
        if (result_dir / filename).is_file()
    ]
    if not available:
        raise ToolError(f'no CSV files in {result_dir}')

    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    output = result_dir / 'combined_384_plate.xlsx'
    tmp_output = output.with_suffix('.tmp.xlsx')
    try:
        frames = [(sheet, pd.read_csv(path)) for sheet, path in available]
        with pd.ExcelWriter(tmp_output, engine='openpyxl') as writer:
            for sheet, frame in frames:
                frame.to_excel(writer, sheet_name=sheet, index=False)

        workbook = load_workbook(tmp_output)
        for sheet in workbook.worksheets:
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column in sheet.columns:
                width = max(
                    len(str(cell.value if cell.value is not None else ''))
                    for cell in column
                ) + 2
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width, 40)
        workbook.save(tmp_output)
        tmp_output.replace(output)
    except ToolError:
        tmp_output.unlink(missing_ok=True)
        raise
    except Exception as exc:
        tmp_output.unlink(missing_ok=True)
        detail = _first_line(exc)
        raise ToolError(f'cannot export {result_dir.name}: {detail}') from exc

    return output
