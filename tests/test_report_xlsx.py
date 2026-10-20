import pytest
from openpyxl import load_workbook

import report_xlsx
import tools


def test_exports_result_folder(tmp_path):
    result_dir = tmp_path / 'Результат 1'
    result_dir.mkdir()
    (result_dir / 'combined_384_plate.csv').write_text(
        'A,B\n1,2\n', encoding='utf-8'
    )
    (result_dir / 'metadata.csv').write_text(
        'sample,value\nsoy1,10\n', encoding='utf-8'
    )

    output = report_xlsx.export_workbook(result_dir)

    assert output.name == 'combined_384_plate.xlsx'
    assert output.is_file()
    workbook = load_workbook(output)
    assert workbook.sheetnames == ['combined', 'metadata']
    assert workbook['combined']['A1'].value == 'A'
    assert workbook['combined']['B2'].value == 2
    assert workbook['combined'].freeze_panes == 'A2'
    assert workbook['combined']['A1'].font.bold is True


def test_missing_folder_is_rejected(tmp_path):
    with pytest.raises(tools.ToolError, match='folder not found'):
        report_xlsx.export_workbook(tmp_path / 'missing')


def test_folder_without_csv_is_rejected(tmp_path):
    result_dir = tmp_path / 'empty'
    result_dir.mkdir()

    with pytest.raises(tools.ToolError, match='no CSV files'):
        report_xlsx.export_workbook(result_dir)
