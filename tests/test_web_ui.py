import json
import os
import shutil
import tempfile
from pathlib import Path

import pytest
from nicegui.testing import User
from openpyxl import Workbook

pytest_plugins = ['nicegui.testing.user_plugin']

WORKBOOK = Path(__file__).resolve().parents[1] / 'src' / 'планшеты соя пробелы.xlsx'

BASE = Path(tempfile.mkdtemp(prefix='avatar-web-'))
shutil.copy(WORKBOOK, BASE / 'plates.xlsx')
os.environ['AVATAR_WEB_BASE'] = str(BASE)


def _workbook(path, names):
    workbook = Workbook()
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        for column in range(2, 14):
            sheet.cell(row=2, column=column).value = column - 1
        for index, row in enumerate(range(3, 11)):
            sheet.cell(row=row, column=1).value = chr(ord('A') + index)
        sheet.cell(row=3, column=2).value = 'sample'
    workbook.save(path)


_workbook(BASE / 'zz_other.xlsx', ['o1', 'o2', 'o3', 'o4'])

pytestmark = pytest.mark.nicegui_main_file('tests/web_main.py')


async def test_home_shows_workbook(user: User):
    await user.open('/')

    await user.should_see('AVATAR COMBINER')
    await user.should_see('plates.xlsx')


async def test_switching_workbook_refreshes_sheets_and_status(user: User):
    await user.open('/')
    other = str(BASE / 'zz_other.xlsx')
    select = next(iter(user.find(marker='workbook').elements))
    before = set(BASE.glob('Результат *'))

    with user.client:
        select._handle_value_change(other)

    sheet = next(iter(user.find(marker='sheet_1').elements))
    assert list(sheet.options)[:4] == ['o1', 'o2', 'o3', 'o4']
    status = next(iter(user.find(marker='selection_status').elements))
    assert status.text == '4 sheets selected'
    info = next(iter(user.find(marker='workbook_info').elements))
    assert 'o1' in info.text

    user.find('Run').click()
    await user.should_see('Done:', retries=50)

    created = set(BASE.glob('Результат *')) - before
    assert len(created) == 1
    manifest = json.loads((next(iter(created)) / 'manifest.json').read_text('utf-8'))
    assert manifest['excel_file'] == 'zz_other.xlsx'


async def test_home_shows_workbook_summary(user: User):
    await user.open('/')

    await user.should_see('sheets;')
    await user.should_see('пл1')


async def test_sheet_selects_are_populated_from_workbook(user: User):
    await user.open('/')

    select = next(iter(user.find(marker='sheet_1').elements))

    options = list(select.options)

    assert options[:4] == ['пл1', 'пл2', 'пл3', 'пл4']
    assert len(options) >= 5
    assert 'пл5' in options


async def test_home_shows_valid_selection(user: User):
    await user.open('/')

    await user.should_see('4 sheets selected')


async def test_nav_lists_every_tool(user: User):
    await user.open('/')

    for name in (
        'Pipeline',
        'Results',
        'Primer Tm',
        'Primer pair',
        'DNA',
        'Units',
        'Structure',
        'PCR',
        'Import sequence',
        'Plate layout',
        'XLSX export',
    ):
        await user.should_see(name)


async def test_tools_page_lists_tools(user: User):
    await user.open('/tools')

    for name in (
        'Primer Tm',
        'Primer pair',
        'DNA',
        'Units',
        'Structure',
        'PCR',
        'Import sequence',
        'Plate layout',
        'XLSX export',
    ):
        await user.should_see(name)


async def test_tool_query_selects_tab(user: User):
    await user.open('/tools?tool=pcr')

    await user.should_see('Simulate')


async def test_tm_tool_produces_output(user: User):
    await user.open('/tools?tool=tm')

    user.find(marker='tm_sequence').type('ACGTACGTACGT')
    user.find(marker='tm_button').click()

    await user.should_see('kcal/mol')


async def test_units_tool_produces_output(user: User):
    await user.open('/tools?tool=units')

    user.find(marker='unit_value').type('25 uM')
    user.find(marker='unit_target').type('nM')
    user.find(marker='unit_button').click()

    await user.should_see('25000 nM')


async def test_language_switch(user: User):
    await user.open('/')
    user.find('RU').click()

    await user.open('/')
    await user.should_see('Пайплайн')
    await user.should_see('Tm праймера')
    await user.should_see('Запустить')

    user.find('EN').click()
    await user.open('/')
    await user.should_see('Pipeline')


async def test_run_from_web(user: User):
    await user.open('/')
    await user.should_see('Run')

    user.find('Run').click()

    await user.should_see('Done:', retries=50)
