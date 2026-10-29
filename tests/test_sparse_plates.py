"""AV-9: fewer than four plates and blank / NaN wells."""

import math
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

import app_config
import combine_plates
import extract_plates_from_excel as extract
import run_pipeline

WORKBOOK = Path(__file__).resolve().parents[1] / 'src' / 'планшеты соя пробелы.xlsx'


def _labeled_workbook(path, names, blank_wells=None, nan_wells=None):
    blank_wells = blank_wells or set()
    nan_wells = nan_wells or set()
    workbook = Workbook()
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        for column in range(2, 14):
            sheet.cell(row=2, column=column).value = column - 1
        for index, row in enumerate(range(3, 11)):
            sheet.cell(row=row, column=1).value = chr(ord('A') + index)
            for column in range(2, 14):
                key = (sheet.title, row, column)
                if key in nan_wells:
                    sheet.cell(row=row, column=column).value = float('nan')
                elif key in blank_wells:
                    sheet.cell(row=row, column=column).value = None
                else:
                    sheet.cell(row=row, column=column).value = (
                        f'{sheet.title}-R{row}C{column}'
                    )
    workbook.save(path)


def test_sample_workbook_blanks_survive_pipeline(tmp_path):
    shutil.copy(WORKBOOK, tmp_path / 'plates.xlsx')

    result = run_pipeline.run(base_dir=tmp_path, log=lambda message: None)

    ids = (result / 'combined_384_plate_ids.txt').read_text(encoding='utf-8').splitlines()
    assert len(ids) == 384
    assert sum(1 for value in ids if value == '') == 16
    assert all(value.lower() != 'nan' for value in ids)


def test_two_plates_pad_empty_384_slots(tmp_path):
    path = tmp_path / 'two.xlsx'
    _labeled_workbook(path, ['p1', 'p2', 'extra'])

    config = app_config.build_config(tmp_path, excel_file=path, sheets=['p1', 'p2'])
    assert list(config['sheets']) == ['plate1', 'plate2']

    result = run_pipeline.run(base_dir=tmp_path, config=config, log=lambda message: None)

    ids = (result / 'combined_384_plate_ids.txt').read_text(encoding='utf-8').splitlines()
    assert len(ids) == 384
    # Plates 3–4 are empty pads (192 wells); plates 1–2 are full here.
    assert sum(1 for value in ids if value == '') == 192
    assert (result / 'plate3.csv').is_file()
    assert (result / 'plate4.csv').is_file()
    plate3 = (result / 'plate3.csv').read_text(encoding='utf-8')
    assert plate3.count('\n') >= 96
    assert all(
        line.split(',')[0] == '' or line.startswith('id')
        for line in plate3.splitlines()
    )


def test_one_plate_still_builds_384(tmp_path):
    path = tmp_path / 'one.xlsx'
    _labeled_workbook(path, ['only', 'spare'])

    config = app_config.build_config(tmp_path, excel_file=path, sheets=['only'])
    result = run_pipeline.run(base_dir=tmp_path, config=config, log=lambda message: None)

    ids = (result / 'combined_384_plate_ids.txt').read_text(encoding='utf-8').splitlines()
    assert len(ids) == 384
    assert sum(1 for value in ids if value) == 96


def test_nan_cells_become_empty_ids(tmp_path):
    path = tmp_path / 'nan.xlsx'
    _labeled_workbook(
        path,
        ['a', 'b', 'c', 'd'],
        nan_wells={('a', 3, 2), ('b', 4, 3)},
        blank_wells={('c', 5, 4)},
    )
    config = app_config.build_config(
        tmp_path, excel_file=path, sheets=['a', 'b', 'c', 'd']
    )
    result = run_pipeline.run(base_dir=tmp_path, config=config, log=lambda message: None)
    plate1 = list(
        line.split(',')[0]
        for line in (result / 'plate1.csv').read_text(encoding='utf-8').splitlines()[1:]
    )
    assert '' in plate1
    assert 'nan' not in {value.lower() for value in plate1}


def test_build_config_rejects_empty_sheet_list(tmp_path):
    path = tmp_path / 'x.xlsx'
    _labeled_workbook(path, ['a', 'b', 'c', 'd'])

    with pytest.raises(app_config.ConfigError, match='at least one'):
        app_config.build_config(tmp_path, excel_file=path, sheets=[])


def test_combine_pads_missing_plate_files(tmp_path):
    (tmp_path / 'plate1.csv').write_text(
        'id,row,column\nS1,A,1\n', encoding='utf-8'
    )
    combine_plates.main(output_dir=tmp_path, input_dir=tmp_path, log=lambda m: None)
    ids = (tmp_path / 'combined_384_plate_ids.txt').read_text(encoding='utf-8').splitlines()
    assert len(ids) == 384
    assert ids.count('S1') == 1


def test_cell_to_id_normalizes_nan():
    assert extract.cell_to_id(None) == ''
    assert extract.cell_to_id(float('nan')) == ''
    assert extract.cell_to_id('nan') == ''
    assert extract.cell_to_id('  sample  ') == 'sample'
    assert math.isnan(float('nan'))
