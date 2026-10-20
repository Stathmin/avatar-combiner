"""Browser interaction matrix for the local web UI.

Run explicitly:

    uv run pytest dev/test_web_modes.py -q

Starts the real NiceGUI app on a temporary folder and drives Chromium through
every user-facing mode: validation, uploads, pipeline run, results, all nine
tools, navigation and language switching.
"""

from __future__ import annotations

import itertools
import json
import os
import signal
import subprocess
import sys
import time
from contextlib import contextmanager
from pathlib import Path
from urllib.request import urlopen

import pytest
from openpyxl import Workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
PORTS = itertools.count(8160)

PCR_TEMPLATE = (
    'GCTAAAGACAATTACATAACATACACGTCAGCACGAAACTTGTTGGCCCAGTGTGAATCGCTTAAGGG'
    'TTAAGTAAGTGTGATGCATACGCCTTTACTTGCTGTGTCCACCCCATCGGACTGGCATTTTTATTACA'
    'CTCAGAAACAGAACTCGGGTAATTTTGACAGGTCACGCAGAGGCGCGCCCTCCTGAAGTGCGTG'
)
from Bio.Seq import Seq  # noqa: E402

PCR_FORWARD = PCR_TEMPLATE[:20]
PCR_REVERSE = str(Seq(PCR_TEMPLATE[-20:]).reverse_complement())


def make_workbook(path, names):
    workbook = Workbook()
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        sheet.cell(row=1, column=1).value = f'{sheet.title} metadata'
        for column in range(2, 14):
            sheet.cell(row=2, column=column).value = column - 1
        for index, row in enumerate(range(3, 11)):
            sheet.cell(row=row, column=1).value = chr(ord('A') + index)
            for column in range(2, 14):
                sheet.cell(row=row, column=column).value = f'{sheet.title}-{row}-{column}'
    workbook.save(path)


def make_jumbled(path):
    workbook = Workbook()
    workbook.active.title = 'one'
    for name in ('two', 'three', 'four'):
        workbook.create_sheet(name)
    workbook['one'].cell(row=10, column=13).value = 'stray'
    workbook.save(path)


def prepare_env(base: Path, seed_result: bool = False):
    base.mkdir(parents=True, exist_ok=True)
    make_workbook(base / 'base.xlsx', ['p1', 'p2', 'p3', 'p4', 'p5'])
    make_workbook(base / 'other.xlsx', ['o1', 'o2', 'o3', 'o4'])
    make_workbook(base / 'extra.xlsx', ['u1', 'u2', 'u3', 'u4'])
    make_jumbled(base / 'jumbled.xlsx')
    (base / 'notes.docx').write_bytes(b'not a workbook')
    (base / 'seq.fa').write_text('>seq1\nACGTACGT\n', encoding='utf-8')
    (base / 'plate.toml').write_text(
        '[well.A1]\nsample = "soy1"\n\n[well.B1]\nsample = "soy2"\n',
        encoding='utf-8',
    )
    alt = base / 'alt'
    alt.mkdir(exist_ok=True)
    make_workbook(alt / 'base.xlsx', ['c1', 'c2', 'c3', 'c4'])
    if seed_result:
        result = base / 'Результат 1'
        result.mkdir()
        (result / 'combined_384_plate.csv').write_text('A,B\n1,2\n', encoding='utf-8')
        (result / 'manifest.json').write_text(
            json.dumps({'excel_file': 'base.xlsx'}), encoding='utf-8'
        )


def wait_for_server(url, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with urlopen(url, timeout=1) as response:
                if response.status == 200:
                    return
        except Exception:
            time.sleep(0.2)
    raise RuntimeError(f'server did not start: {url}')


@contextmanager
def running_web(base: Path):
    port = next(PORTS)
    url = f'http://127.0.0.1:{port}'
    log = open(base / 'server.log', 'w', encoding='utf-8')
    env = dict(os.environ)
    env.pop('PYTEST_CURRENT_TEST', None)
    env.pop('NICEGUI_SCREEN_TEST_PORT', None)
    process = subprocess.Popen(
        [sys.executable, str(ROOT / 'dev' / 'serve_web.py'), str(base), str(port)],
        stdout=log,
        stderr=log,
        env=env,
        start_new_session=True,
    )
    try:
        wait_for_server(url)
        yield url
    finally:
        os.killpg(process.pid, signal.SIGTERM)
        log.close()


@pytest.fixture(scope='session')
def browser():
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            executable_path='/usr/bin/chromium',
            headless=True,
            args=['--no-sandbox'],
        )
        yield browser
        browser.close()


@pytest.fixture
def page(browser):
    context = browser.new_context(
        viewport={'width': 1400, 'height': 1100},
        accept_downloads=True,
    )
    page = context.new_page()
    page.set_default_timeout(5000)
    yield page
    context.close()


@pytest.fixture(scope='module')
def web(tmp_path_factory):
    base = tmp_path_factory.mktemp('web')
    prepare_env(base)
    with running_web(base) as url:
        yield base, url


@pytest.fixture(scope='module')
def web_seed(tmp_path_factory):
    base = tmp_path_factory.mktemp('seed')
    prepare_env(base, seed_result=True)
    with running_web(base) as url:
        yield base, url


@pytest.fixture(scope='module')
def web_empty(tmp_path_factory):
    base = tmp_path_factory.mktemp('empty')
    with running_web(base) as url:
        yield base, url


def open_home(page, url):
    page.goto(url)
    expect(page.get_by_text('4 sheets selected')).to_be_visible()


def workbook_select(page):
    return page.locator('.q-select').nth(0)


def sheet_selects(page):
    return [page.locator('.q-select').nth(index) for index in range(1, 5)]


def choose(page, select, option):
    select.click()
    page.locator('.q-menu .q-item').filter(has_text=option).first.click()


def visible_panel(page):
    return page.locator('.q-tab-panel:visible')


def tool_output(page):
    expect(visible_panel(page).locator('pre')).to_be_visible()
    return visible_panel(page).locator('pre')


def test_home_initial_state(web, page):
    base, url = web
    open_home(page, url)

    expect(workbook_select(page)).to_contain_text('base.xlsx')
    expect(page.get_by_text('4 sheets selected')).to_be_visible()
    expect(page.get_by_text('10×13').first).to_be_visible()
    expect(page.get_by_role('button', name='Run')).to_be_enabled()
    assert page.get_by_text('Sheets', exact=True).count() == 1
    assert page.locator('.q-uploader').count() == 1


def test_duplicate_sheet_disables_run(web, page):
    _, url = web
    open_home(page, url)

    choose(page, sheet_selects(page)[1], 'p1')
    expect(page.get_by_text('unique')).to_be_visible()
    expect(page.get_by_role('button', name='Run')).to_be_disabled()

    choose(page, sheet_selects(page)[1], 'p2')
    expect(page.get_by_text('4 sheets selected')).to_be_visible()
    expect(page.get_by_role('button', name='Run')).to_be_enabled()


def test_invalid_range_disables_run(web, page):
    _, url = web
    open_home(page, url)

    page.get_by_label('Data range').fill('B3')
    expect(page.get_by_text('use a range like B3:M10')).to_be_visible()
    expect(page.get_by_role('button', name='Run')).to_be_disabled()

    page.get_by_label('Data range').fill('B3:M10')
    expect(page.get_by_role('button', name='Run')).to_be_enabled()


def test_workbook_switch_updates_sheets_and_manifest(web, page):
    base, url = web
    open_home(page, url)

    choose(page, workbook_select(page), 'other.xlsx')
    expect(sheet_selects(page)[3]).to_contain_text('o4')
    expect(page.get_by_text('4 sheets selected')).to_be_visible()

    before = set(base.glob('Результат *'))
    page.get_by_role('button', name='Run').click()
    expect(page.get_by_text('Done:')).to_be_visible()
    created = set(base.glob('Результат *')) - before
    manifest = json.loads((next(iter(created)) / 'manifest.json').read_text('utf-8'))
    assert manifest['excel_file'] == 'other.xlsx'


def test_upload_valid_replaces_selection(web, page):
    base, url = web
    open_home(page, url)
    source = base.parent / 'extra-source.xlsx'
    make_workbook(source, ['u1', 'u2', 'u3', 'u4'])

    page.locator('input[type=file]').first.set_input_files(str(source))
    expect(workbook_select(page)).to_contain_text('extra-source.xlsx')
    expect(sheet_selects(page)[0]).to_contain_text('u1')
    expect(page.get_by_text('4 sheets selected')).to_be_visible()
    expect(page.locator('.q-uploader__file')).to_have_count(0)


def test_upload_same_name_is_renamed(web, page):
    base, url = web
    open_home(page, url)

    page.locator('input[type=file]').first.set_input_files(
        str(base / 'alt' / 'base.xlsx')
    )
    expect(workbook_select(page)).to_contain_text('base-uploaded.xlsx')
    expect(sheet_selects(page)[0]).to_contain_text('c1')


def test_upload_jumbled_is_rejected(web, page):
    base, url = web
    open_home(page, url)

    page.locator('input[type=file]').first.set_input_files(str(base / 'jumbled.xlsx'))
    expect(page.get_by_text('File rejected')).to_be_visible()
    expect(workbook_select(page)).to_contain_text('base.xlsx')
    assert not any(
        'jumbled' in path.name for path in (base / '.uploads').glob('*.xlsx')
    )


def test_upload_docx_is_rejected(web, page):
    base, url = web
    open_home(page, url)

    page.locator('input[type=file]').first.set_input_files(str(base / 'notes.docx'))
    expect(page.get_by_text('Only .xlsx or .xlsm')).to_be_visible()
    assert not (base / '.uploads' / 'notes.docx').exists()


def test_run_logs_result_and_download(web, page):
    base, url = web
    open_home(page, url)

    page.get_by_role('button', name='Run').click()
    expect(page.get_by_text('Done:')).to_be_visible()
    expect(page.get_by_text('Пайплайн выполнен успешно!')).to_be_visible()
    expect(page.get_by_text('manifest.json').first).to_be_visible()

    download = page.get_by_role('button', name='combined_384_plate.csv')
    with page.expect_download() as info:
        download.click()
    assert info.value.suggested_filename == 'combined_384_plate.csv'


def test_results_empty(web_empty, page):
    _, url = web_empty
    page.goto(url + '/results')
    expect(page.get_by_text('No results yet.')).to_be_visible()


def test_results_list_and_download(web_seed, page):
    base, url = web_seed
    page.goto(url + '/results')

    expect(page.get_by_text('Результат 1')).to_be_visible()
    with page.expect_download() as info:
        page.get_by_role('button', name='combined_384_plate.csv').click()
    assert info.value.suggested_filename == 'combined_384_plate.csv'


TOOL_TABS = (
    ('tm', 'Primer Tm'),
    ('pair', 'Primer pair'),
    ('dna', 'DNA'),
    ('units', 'Units'),
    ('structure', 'Structure'),
    ('pcr', 'PCR'),
    ('import', 'Files'),
    ('layout', 'Files'),
    ('xlsx', 'Files'),
)


def test_navigation_selects_every_tool(web, page):
    _, url = web
    open_home(page, url)

    for key, tab_name in TOOL_TABS:
        page.goto(f'{url}/tools?tool={key}')
        expect(page.locator('.q-tab--active')).to_contain_text(tab_name)
        expect(page.locator('.q-tab-panel:visible')).to_be_visible()


def test_tool_primer_tm(web, page):
    _, url = web
    page.goto(url + '/tools?tool=tm')

    panel = visible_panel(page)
    panel.get_by_label('Sequence').fill('ACGTACGTACGT')
    panel.get_by_role('button', name='Analyse').click()
    expect(tool_output(page)).to_contain_text('kcal/mol')
    expect(tool_output(page)).to_contain_text('Tm')

    panel.get_by_label('Sequence').fill('XXXX')
    panel.get_by_role('button', name='Analyse').click()
    expect(tool_output(page)).to_contain_text('Error:')


def test_tool_pair(web, page):
    _, url = web
    page.goto(url + '/tools?tool=pair')

    panel = visible_panel(page)
    panel.get_by_label('Forward').fill('GTAAAACGACGGCCAGT')
    panel.get_by_label('Reverse').fill('GCGTACGTAGCTAGCTA')
    panel.get_by_role('button', name='Analyse pair').click()
    expect(tool_output(page)).to_contain_text('ΔTm')
    expect(tool_output(page)).to_contain_text('Heterodimer')


def test_tool_dna(web, page):
    _, url = web
    page.goto(url + '/tools?tool=dna')

    panel = visible_panel(page)
    panel.get_by_label('Sequence').fill('ATGC')
    panel.get_by_role('button', name='Transform').click()
    expect(tool_output(page)).to_contain_text('TACG')
    expect(tool_output(page)).to_contain_text('GCAT')


def test_tool_units(web, page):
    _, url = web
    page.goto(url + '/tools?tool=units')

    panel = visible_panel(page)
    panel.get_by_label('Value').fill('25 uM')
    panel.get_by_label('Target unit').fill('nM')
    panel.get_by_role('button', name='Convert').click()
    expect(tool_output(page)).to_contain_text('25000 nM')


def test_tool_structure(web, page):
    _, url = web
    page.goto(url + '/tools?tool=structure')

    panel = visible_panel(page)
    panel.get_by_label('Sequence').fill('ACGTACGTACGT')
    panel.get_by_role('button', name='Fold').click()
    expect(tool_output(page)).to_contain_text('((((....))))')


def test_tool_pcr(web, page):
    _, url = web
    page.goto(url + '/tools?tool=pcr')

    panel = visible_panel(page)
    panel.get_by_label('Template').fill(PCR_TEMPLATE)
    panel.get_by_label('Forward').fill(PCR_FORWARD)
    panel.get_by_label('Reverse').fill(PCR_REVERSE)
    panel.get_by_role('button', name='Simulate').click()
    expect(tool_output(page)).to_contain_text('200 nt')


def test_tool_import_fasta(web, page):
    base, url = web
    page.goto(url + '/tools?tool=import')

    panel = visible_panel(page)
    panel.locator('input[type=file]').nth(0).set_input_files(str(base / 'seq.fa'))
    expect(tool_output(page)).to_contain_text('fasta')
    expect(tool_output(page)).to_contain_text('8 nt')
    expect(tool_output(page)).to_contain_text('seq1')


def test_tool_layout(web, page):
    base, url = web
    page.goto(url + '/tools?tool=layout')

    panel = visible_panel(page)
    panel.locator('input[type=file]').nth(1).set_input_files(str(base / 'plate.toml'))
    expect(tool_output(page)).to_contain_text('sample')
    expect(tool_output(page)).to_contain_text('2')


def test_tool_xlsx_export(web_seed, page):
    base, url = web_seed
    page.goto(url + '/tools?tool=xlsx')

    panel = visible_panel(page)
    choose(page, panel.locator('.q-select'), 'Результат 1')
    panel.get_by_role('button', name='Export XLSX').click()
    expect(tool_output(page)).to_contain_text('combined_384_plate.xlsx')
    assert (base / 'Результат 1' / 'combined_384_plate.xlsx').is_file()


def test_language_switch_preserves_selection(web, page):
    _, url = web
    open_home(page, url)

    page.get_by_role('button', name='RU', exact=True).click()
    expect(page.get_by_text('Пайплайн').first).to_be_visible()
    expect(page.get_by_text('Выбрано 4 листа')).to_be_visible()
    expect(workbook_select(page)).to_contain_text('base.xlsx')

    page.get_by_role('button', name='EN', exact=True).click()
    expect(page.get_by_text('Pipeline').first).to_be_visible()


def test_tools_page_has_no_duplicate_index(web, page):
    _, url = web
    page.goto(url + '/tools')

    assert page.locator('.tool-index').count() == 0
    assert page.locator('.q-tab-panel:visible').count() == 1
    assert page.get_by_text('All tools').count() == 0


def test_sheet_choice_is_used_in_manifest(web, page):
    base, url = web
    open_home(page, url)

    choose(page, sheet_selects(page)[0], 'p5')
    expect(page.get_by_text('4 sheets selected')).to_be_visible()

    before = set(base.glob('Результат *'))
    page.get_by_role('button', name='Run').click()
    expect(page.get_by_text('Done:')).to_be_visible()
    created = set(base.glob('Результат *')) - before
    manifest = json.loads((next(iter(created)) / 'manifest.json').read_text('utf-8'))
    assert manifest['sheets']['plate1']['sheet_name'] == 'p5'


def test_tools_errors_are_informative(web, page):
    base, url = web
    page.goto(url + '/tools?tool=units')
    panel = visible_panel(page)
    panel.get_by_label('Value').fill('25 xyz')
    panel.get_by_label('Target unit').fill('nM')
    panel.get_by_role('button', name='Convert').click()
    expect(tool_output(page)).to_contain_text('Error:')

    page.goto(url + '/tools?tool=pcr')
    panel = visible_panel(page)
    panel.get_by_label('Template').fill(PCR_TEMPLATE)
    panel.get_by_label('Forward').fill('A' * 20)
    panel.get_by_label('Reverse').fill('T' * 20)
    panel.get_by_role('button', name='Simulate').click()
    expect(tool_output(page)).to_contain_text('no PCR product')

    page.goto(url + '/tools?tool=import')
    panel = visible_panel(page)
    panel.locator('input[type=file]').nth(0).set_input_files(str(base / 'notes.docx'))
    expect(tool_output(page)).to_contain_text('Error:')


def test_tabs_and_drawer_have_unique_names(web, page):
    _, url = web
    page.goto(url + '/tools')

    tabs = [text.strip() for text in page.locator('.q-tab').all_inner_texts()]
    assert len(tabs) == len(set(tabs)) == 7

    links = [
        text.strip()
        for text in page.locator('.q-drawer a').all_inner_texts()
    ]
    assert len(links) == len(set(links)) == 11


def test_results_manifest_download(web_seed, page):
    _, url = web_seed
    page.goto(url + '/results')

    with page.expect_download() as info:
        page.get_by_role('button', name='manifest.json').click()
    assert info.value.suggested_filename == 'manifest.json'


def test_language_switch_on_tools(web, page):
    _, url = web
    page.goto(url + '/tools')

    page.get_by_role('button', name='RU', exact=True).click()
    expect(page.get_by_text('Tm праймера').first).to_be_visible()
    expect(page.get_by_text('ПЦР').first).to_be_visible()

    page.get_by_role('button', name='EN', exact=True).click()
    expect(page.get_by_text('Primer Tm').first).to_be_visible()


def test_no_workbook_shows_error_and_disables_run(web_empty, page):
    _, url = web_empty
    page.goto(url)
    expect(page.get_by_text('no Excel file', exact=False)).to_be_visible()
    expect(page.get_by_role('button', name='Run')).to_be_disabled()
