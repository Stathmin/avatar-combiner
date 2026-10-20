"""Browser end-to-end check for the web upload flow.

Run from the repo root:

    uv run python dev/web_e2e.py

Uses the system Chromium through Playwright; no browser download required.
"""

from __future__ import annotations

import json
import os
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib.request import urlopen

from openpyxl import Workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
PORT = 8150
URL = f'http://127.0.0.1:{PORT}/'


def make_workbook(path, names):
    workbook = Workbook()
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        sheet.cell(row=1, column=1).value = f'{sheet.title} metadata'
        for column in range(2, 14):
            sheet.cell(row=2, column=column).value = column - 1
        for index, row in enumerate(range(3, 11)):
            sheet.cell(row=row, column=1).value = chr(ord('A') + index)
            for column in range(2, 14):
                sheet.cell(row=row, column=column).value = f'{sheet.title}-{row}-{column}'
    workbook.save(path)


def make_jumbled(path):
    workbook = Workbook()
    workbook.active.title = 'one'
    for name in ('two', 'three', 'four'):
        workbook.create_sheet(name)
    workbook['one'].cell(row=10, column=13).value = 'stray'
    workbook.save(path)


def wait_for_server(timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with urlopen(URL, timeout=1) as response:
                if response.status == 200:
                    return
        except Exception:
            time.sleep(0.3)
    raise RuntimeError('web server did not start')


def main():
    base = Path(tempfile.mkdtemp(prefix='avatar-e2e-'))
    make_workbook(base / 'base.xlsx', ['p1', 'p2', 'p3', 'p4'])
    make_workbook(base / 'biometria.xlsx', ['b1', 'b2', 'b3', 'b4'])
    make_jumbled(base / 'jumbled.xlsx')
    (base / 'notes.docx').write_bytes(b'not a workbook')

    log = open(base / 'server.log', 'w', encoding='utf-8')
    server = subprocess.Popen(
        [sys.executable, str(ROOT / 'dev' / 'serve_web.py'), str(base), str(PORT)],
        stdout=log,
        stderr=log,
        start_new_session=True,
    )
    try:
        wait_for_server()
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(
                executable_path='/usr/bin/chromium',
                headless=True,
                args=['--no-sandbox'],
            )
            page = browser.new_page(viewport={'width': 1400, 'height': 1100})
            page.goto(URL)
            page.wait_for_selector('text=4 sheets selected', timeout=20000)
            artifacts = ROOT / 'dev' / 'artifacts'
            artifacts.mkdir(exist_ok=True)

            print('initial workbook:', page.locator('.q-field').first.inner_text())
            page.screenshot(path=str(artifacts / '01-initial.png'), full_page=True)

            upload = page.locator('input[type=file]').first
            assert upload.get_attribute('accept') == '.xlsx,.xlsm'

            upload.set_input_files(str(base / 'jumbled.xlsx'))
            page.wait_for_timeout(1500)
            notices = page.locator('.q-notification').all_inner_texts()
            print('jumbled notices:', notices)
            assert any('File rejected' in notice for notice in notices)
            assert [path.name for path in (base / '.uploads').glob('*')] == []
            page.screenshot(path=str(artifacts / '02-jumbled.png'), full_page=True)

            upload.set_input_files(str(base / 'biometria.xlsx'))
            page.wait_for_timeout(2500)
            selects = page.locator('.q-select').all_inner_texts()[:5]
            print('selects:', selects)
            assert 'biometria-uploaded.xlsx' in selects[0]
            assert all(f'b{index}' in selects[index] for index in range(1, 5))
            assert page.locator('text=4 sheets selected').count() == 1
            assert page.locator('text=b1 10×13').count() >= 1
            page.screenshot(path=str(artifacts / '03-uploaded.png'), full_page=True)

            before = set(base.glob('Результат *'))
            page.get_by_role('button', name='Run').click()
            page.wait_for_selector('text=Done:', timeout=30000)
            created = set(base.glob('Результат *')) - before
            assert len(created) == 1, created
            manifest = json.loads((next(iter(created)) / 'manifest.json').read_text('utf-8'))
            print('manifest workbook:', manifest['excel_file'])
            assert manifest['excel_file'].endswith('biometria-uploaded.xlsx'), manifest['excel_file']
            page.screenshot(path=str(artifacts / '04-result.png'), full_page=True)

            upload.set_input_files(str(base / 'notes.docx'))
            page.wait_for_timeout(1200)
            notices = page.locator('.q-notification').all_inner_texts()
            print('docx notices:', notices)
            assert any('.xlsx' in notice and 'xlsm' in notice for notice in notices)
            assert not (base / '.uploads' / 'notes.docx').exists()
            browser.close()
    finally:
        os.killpg(server.pid, signal.SIGTERM)
        log.close()


if __name__ == '__main__':
    main()
